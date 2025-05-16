"""Summary sheet generator."""

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList

from va_tool.reporting.sheets.base import BaseSheetGenerator
from va_tool.utils import write_df_to_sheet


class SummarySheetGenerator(BaseSheetGenerator):
    """Generator for the Summary sheet."""
    
    def __init__(self):
        """Initialize the generator."""
        super().__init__(title="Summary")
    
    def generate(self, wb, results_data, original_df=None, processed_df=None, **kwargs):
        """
        Generate the Summary sheet.
        
        Args:
            wb: Excel workbook
            results_data: Dictionary with analysis results
            original_df: Original vulnerability DataFrame
            processed_df: Processed vulnerability DataFrame
            **kwargs: Additional arguments
        
        Returns:
            The worksheet
        """
        self.logger.info("Generating Summary sheet")
        ws = super().generate(wb)
        
        # Add title
        self.add_title(
            ws, "Vulnerability Assessment Summary", 
            font_size=16, merge_range='A1:H1'
        )
        
        # Add risk counts section
        self.add_section_title(ws, "Risk Level Distribution", cell="A3")
        self.add_risk_comparison(ws, original_df, processed_df)
        
        # Add Top 10 CVEs section
        self.add_section_title(ws, "Top 10 Most Common CVEs", cell="A12")
        self.add_top_cves(ws, results_data)
        
        # Add bucket summary section
        self.add_section_title(ws, "Vulnerability Distribution by Type", cell="A25")
        self.add_bucket_summary(ws, results_data)
        
        return ws
    
    def add_risk_comparison(self, ws, original_df, processed_df):
        """Add risk comparison table and chart."""
        # Create risk comparison table
        if original_df is not None:
            before_risk_counts = original_df["Risk"].value_counts().reset_index()
            before_risk_counts.columns = ["Risk Level", "Before Count"]
        else:
            before_risk_counts = pd.DataFrame({
                "Risk Level": ["Critical", "High", "Medium", "Low"],
                "Before Count": [0, 0, 0, 0]
            })
        
        if processed_df is not None:
            after_risk_counts = processed_df["Risk"].value_counts().reset_index()
            after_risk_counts.columns = ["Risk Level", "After Count"]
        else:
            after_risk_counts = pd.DataFrame({
                "Risk Level": ["Critical", "High", "Medium", "Low"],
                "After Count": [0, 0, 0, 0]
            })
        
        # Merge before and after counts
        risk_comparison = pd.merge(
            before_risk_counts, after_risk_counts, 
            on="Risk Level", how="outer"
        ).fillna(0)
        
        # Sort by severity
        severity_order = ["Critical", "High", "Medium", "Low"]
        risk_comparison["Risk Level"] = pd.Categorical(
            risk_comparison["Risk Level"],
            categories=severity_order, ordered=True
        )
        risk_comparison = risk_comparison.sort_values("Risk Level")
        
        # Write table
        write_df_to_sheet(ws, risk_comparison, start_row=4, start_col=1)
        
        # Create chart
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Risk Level Distribution"
        chart.y_axis.title = "Count"
        chart.x_axis.title = "Risk Level"
        
        # Add data to chart
        data = Reference(ws, min_col=2, max_col=3, min_row=4, max_row=8)
        cats = Reference(ws, min_col=1, max_col=1, min_row=5, max_row=8)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 10
        chart.width = 20
        
        # Add data labels
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        
        # Add chart to worksheet
        ws.add_chart(chart, "I2")
    
    def add_top_cves(self, ws, results_data):
        """Add top CVEs table and chart."""
        top_cves = pd.DataFrame(results_data.get("top_cves", []))
        
        if not top_cves.empty and "CVE" in top_cves.columns and "Count" in top_cves.columns:
            # Limit to top 10
            top_cves = top_cves.head(10)
            
            # Write table
            write_df_to_sheet(ws, top_cves, start_row=13, start_col=1)
            
            # Create chart
            cve_chart = BarChart()
            cve_chart.type = "col"
            cve_chart.style = 10
            cve_chart.title = "Top 10 CVEs by Occurrence"
            cve_chart.y_axis.title = "Count"
            cve_chart.x_axis.title = "CVE"
            
            # Add data to chart
            cve_data = Reference(ws, min_col=2, max_col=2, min_row=13, max_row=24)
            cve_cats = Reference(ws, min_col=1, max_col=1, min_row=14, max_row=24)
            cve_chart.add_data(cve_data, titles_from_data=True)
            cve_chart.set_categories(cve_cats)
            cve_chart.height = 15
            cve_chart.width = 20
            
            # Add data labels
            cve_chart.dataLabels = DataLabelList()
            cve_chart.dataLabels.showVal = True
            
            # Add chart to worksheet
            ws.add_chart(cve_chart, "I13")
    
    def add_bucket_summary(self, ws, results_data):
        """Add bucket summary table and chart."""
        bucket_summary = pd.DataFrame(results_data.get("bucket_summary", []))
        
        if not bucket_summary.empty:
            # Sort by count and take top 10
            top_buckets = bucket_summary.sort_values("Count", ascending=False).head(10)
            
            # Write table
            write_df_to_sheet(ws, top_buckets, start_row=26, start_col=1)
            
            # Create pie chart
            pie = PieChart()
            pie.title = "Vulnerability Types"
            
            # Add data to chart
            bucket_data = Reference(ws, min_col=2, max_col=2, min_row=26, max_row=37)
            bucket_labels = Reference(ws, min_col=1, max_col=1, min_row=27, max_row=37)
            pie.add_data(bucket_data, titles_from_data=True)
            pie.set_categories(bucket_labels)
            pie.height = 15
            pie.width = 20
            
            # Add data labels
            pie.dataLabels = DataLabelList()
            pie.dataLabels.showPercent = True
            
            # Add chart to worksheet
            ws.add_chart(pie, "I26")